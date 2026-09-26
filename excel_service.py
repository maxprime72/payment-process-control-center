import re
import io
import datetime
import openpyxl  # type: ignore
import pandas as pd
import numpy as np
from typing import Dict, List, Tuple, Any, Optional, cast

def is_na(val: Any) -> bool:
    """Safe boolean check for NA/None values without triggering array boolean evaluation errors."""
    if val is None:
        return True
    try:
        res = pd.isna(val)
        if isinstance(res, (bool, np.bool_)):
            return bool(res)
        if hasattr(res, "__iter__"):
            return all(res)
        return bool(res)
    except Exception:
        return False

def not_na(val: Any) -> bool:
    """Safe boolean check for non-NA values."""
    return not is_na(val)

# Canonical required and optional column definitions
CANONICAL_COLUMNS = {
    # Required
    "unit": ["unit"],
    "bbu_number": ["bbu number", "bbu no", "bbu_number", "bbu", "bbu nos"],
    "vendor_name": ["vendor name", "vendor", "supplier name", "vendor_name"],
    "po_number": ["po no.", "po no", "po number", "purchase order no", "po_number", "po"],
    "invoice_number": ["vendor invoice no.", "vendor invoice no", "vendor invoice number", "invoice no.", "invoice no", "invoice number"],
    "invoice_date": ["vendor invoice date", "invoice date", "inv date"],
    "grn_number": ["grn no.", "grn no", "grn number", "grn_no", "grn"],
    "grn_date": ["grn date", "grn_date"],
    "gross_invoice_total": ["gross invoice total", "gross invoice amount", "gross total", "total invoice value", "invoice amount"],
    "payment_due_days": ["payment due in days (as per po)", "payment due in days", "payment due days", "payment due in days as per po"],
    "invoice_processing_status": ["invoice processing status", "processing status", "invoice status"],
    "processed": ["processed"],
    "inprocess": ["inprocess", "in process"],
    "total_amount_payable": ["total amount payble", "total amount payable", "amount payable", "payable amount"],

    # Optional
    "item_description": ["bbu description/ item description", "bbu description", "item description", "description"],
    "basic_amount": ["basic amount", "basic value"],
    "ssc_request_number": ["ssc request no.", "ssc request no", "ssc request number", "ssc req no", "ssc number"],
    "ssc_upload_date": ["date - uploaded in ssc", "uploaded in ssc", "ssc upload date", "date uploaded in ssc"],
    "hard_copy_received_date": ["date - hard copies received at lmb office", "hard copies received at lmb office", "hard copy received date"],
    "hard_copy_sent_date": ["date - hard copy sent to f&a", "hard copy sent to f&a", "hard copy sent date"],
    "remark_closure_date": ["remark closure date", "closure date"],
    "kz_number": ["kz no", "kz no.", "kz number"],
    "kz_date": ["kz date"],
    "rt": ["rt"],
    "ld": ["ld"],
    "ka": ["ka"],
    "kd": ["kd"],
    "tds": ["tds"],
}

REQUIRED_KEYS = [
    "unit",
    "bbu_number",
    "vendor_name",
    "po_number",
    "invoice_number",
    "invoice_date",
    "grn_number",
    "grn_date",
    "gross_invoice_total",
    "payment_due_days",
    "invoice_processing_status",
    "processed",
    "inprocess",
    "total_amount_payable",
]

# Specifically excluded from business logic & UI
EXCLUDED_PATTERNS = [
    "pgi", "mdcc", "uom", "transporter", "gatepaas", "gatepass", "delivery challan",
    "dc recieved", "dc sent", "lc no", "remark", "navodaya remark", "to be processed by",
    "soft copy uploded"
]

def clean_column_name(col: Any) -> str:
    """Standardizes string by lowercasing and removing non-alphanumeric characters."""
    if col is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(col).strip().lower())

def is_excluded_column(col_str: str) -> bool:
    low = col_str.lower().strip()
    return any(exc in low for exc in EXCLUDED_PATTERNS)

def match_columns(columns: List[Any]) -> Dict[str, str]:
    """
    Matches raw sheet headers to canonical keys.
    Returns dict of {raw_column_name: canonical_key}
    """
    matched = {}
    used_canonical = set()

    for col in columns:
        if is_excluded_column(str(col)):
            continue
        cleaned = clean_column_name(col)
        if not cleaned:
            continue

        for canonical_key, aliases in CANONICAL_COLUMNS.items():
            if canonical_key in used_canonical:
                continue
            for alias in aliases:
                if clean_column_name(alias) == cleaned:
                    matched[str(col)] = canonical_key
                    used_canonical.add(canonical_key)
                    break
            if str(col) in matched:
                break

    return matched

def detect_main_sheet(file_bytes: bytes) -> Tuple[str, List[str], Dict[str, str]]:
    """
    Examines all sheets in the Excel workbook.
    Returns (best_sheet_name, all_sheet_names, column_mapping_for_best_sheet).
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet_names = wb.sheetnames

    best_sheet = sheet_names[0]
    best_score = -1
    best_mapping = {}

    for name in sheet_names:
        sheet = wb[name]
        # Read header row
        headers = []
        iter_rows_fn = getattr(sheet, "iter_rows", None)
        if callable(iter_rows_fn):
            for row in cast(Any, iter_rows_fn(max_row=5, values_only=True)):
                if any(cell is not None for cell in row):
                    headers = [str(c).strip() for c in row if c is not None]
                    break

        mapping = match_columns(headers)
        # Score based on how many required keys matched
        matched_required = sum(1 for k in REQUIRED_KEYS if k in mapping.values())

        if matched_required > best_score:
            best_score = matched_required
            best_sheet = name
            best_mapping = mapping

    wb.close()
    return best_sheet, sheet_names, best_mapping

def format_date_safe(val: Any) -> Optional[str]:
    """Safely converts various date formats (datetime, Excel serial, string) to YYYY-MM-DD."""
    if is_na(val) or str(val).strip() == "":
        return None

    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")

    # If it's a numeric Excel serial
    if isinstance(val, (int, float)):
        try:
            # Excel base date is 1899-12-30
            dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(val))
            if 1990 <= dt.year <= 2100:
                return dt.strftime("%Y-%m-%d")
        except Exception:
            pass

    # Try string parsing
    s = str(val).strip()
    try:
        dt = pd.to_datetime(s, errors="coerce")
        if not_na(dt):
            return dt.strftime("%Y-%m-%d")
    except Exception:
        pass

    return None

def clean_amount(val: Any) -> float:
    """Safely converts numeric or currency string to float, filtering NaN and Inf."""
    if is_na(val):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val) if not np.isnan(val) and not np.isinf(val) else 0.0

    s = str(val).replace(",", "").replace("₹", "").replace("$", "").strip()
    try:
        f = float(s)
        return 0.0 if np.isnan(f) or np.isinf(f) else f
    except (ValueError, TypeError):
        return 0.0

def preview_excel_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Generates pre-flight validation preview for the upload dialog."""
    sheet_name, sheet_names, mapping = detect_main_sheet(file_bytes)

    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name)
    raw_headers = df.columns.tolist()

    mapping = match_columns(raw_headers)
    found_keys = list(mapping.values())

    required_found = [k for k in REQUIRED_KEYS if k in found_keys]
    required_missing = [k for k in REQUIRED_KEYS if k not in found_keys]
    optional_found = [k for k in found_keys if k not in REQUIRED_KEYS]

    validation_status = "VALID" if len(required_missing) == 0 else "ERROR"
    if validation_status == "ERROR":
        validation_message = f"Missing required columns: {', '.join(required_missing)}"
    else:
        validation_message = f"Successfully validated all required columns in sheet '{sheet_name}'."

    # Invert mapping to read columns
    inv_map = {v: k for k, v in mapping.items()}

    inv_col = inv_map.get("invoice_number")
    blank_invoices = int(float(cast(Any, df[inv_col].isna().sum()))) if inv_col and inv_col in df.columns else 0

    # Duplicates check
    vendor_col = inv_map.get("vendor_name")
    po_col = inv_map.get("po_number")
    duplicate_count = 0
    if inv_col and vendor_col and po_col and all(c in df.columns for c in [inv_col, vendor_col, po_col]):
        dup_series = df.duplicated(subset=[inv_col, vendor_col, po_col], keep=False)
        duplicate_count = int(float(cast(Any, dup_series.sum())))

    sample_rows: List[Dict[str, Any]] = []
    for _, row in df.head(5).iterrows():
        sample_row: Dict[str, Any] = {}
        for raw_col, can_key in mapping.items():
            val = row.get(raw_col)
            sample_row[can_key] = None if is_na(val) else str(val)
        sample_rows.append(sample_row)

    return {
        "filename": filename,
        "sheet_names": sheet_names,
        "detected_sheet": sheet_name,
        "total_rows": len(df),
        "total_columns": len(raw_headers),
        "required_found": required_found,
        "required_missing": required_missing,
        "optional_found": optional_found,
        "sample_rows": sample_rows,
        "validation_status": validation_status,
        "validation_message": validation_message,
        "duplicate_count": duplicate_count,
        "blank_invoice_count": blank_invoices,
    }

def parse_and_normalize_data(file_bytes: bytes, filename: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Parses and normalizes the uploaded file into a clean list of record dicts.
    Validates required columns. Raises ValueError if validation fails.
    """
    sheet_name, sheet_names, mapping = detect_main_sheet(file_bytes)
    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name)

    mapping = match_columns(df.columns.tolist())
    found_keys = set(mapping.values())
    missing_required = [k for k in REQUIRED_KEYS if k not in found_keys]

    if missing_required:
        raise ValueError(
            f"Excel file '{filename}' sheet '{sheet_name}' is missing required columns: {', '.join(missing_required)}"
        )

    inv_map = {v: k for k, v in mapping.items()}

    # Check duplicates by Vendor + PO + Invoice
    inv_col = inv_map.get("invoice_number")
    vendor_col = inv_map.get("vendor_name")
    po_col = inv_map.get("po_number")

    dup_mask = pd.Series(False, index=df.index)
    if inv_col and vendor_col and po_col and all(c in df.columns for c in [inv_col, vendor_col, po_col]):
        dup_mask = df.duplicated(subset=[inv_col, vendor_col, po_col], keep=False)

    records: List[Dict[str, Any]] = []
    for idx, row in df.iterrows():
        rec: Dict[str, Any] = {}

        # Helper to get raw column value safely
        def get_val(key: str) -> Any:
            col = inv_map.get(key)
            if col and col in df.columns:
                v = row.get(col)
                return v if not_na(v) else None
            return None

        # Identification
        rec["unit"] = str(get_val("unit")).strip() if get_val("unit") is not None else None
        bbu = str(get_val("bbu_number")).strip() if get_val("bbu_number") is not None else ""
        rec["bbu_number"] = bbu if bbu else None
        rec["item_description"] = str(get_val("item_description")).strip() if get_val("item_description") is not None else None

        # Non-billable detection
        rec["is_non_billable"] = bool(re.search(r"non-?billable", bbu, re.IGNORECASE))

        # Vendor & PO
        v_name = str(get_val("vendor_name")).strip() if get_val("vendor_name") is not None else None
        rec["vendor_name"] = v_name
        po = str(get_val("po_number")).strip() if get_val("po_number") is not None else None
        rec["po_number"] = po

        # Invoice
        inv_no = str(get_val("invoice_number")).strip() if get_val("invoice_number") is not None else None
        rec["invoice_number"] = inv_no
        rec["invoice_date"] = format_date_safe(get_val("invoice_date"))

        # GRN
        grn_no = str(get_val("grn_number")).strip() if get_val("grn_number") is not None else None
        rec["grn_number"] = grn_no
        rec["grn_date"] = format_date_safe(get_val("grn_date"))

        # Financial amounts
        rec["basic_amount"] = clean_amount(get_val("basic_amount"))
        rec["gross_invoice_total"] = clean_amount(get_val("gross_invoice_total"))
        rec["total_amount_payable"] = clean_amount(get_val("total_amount_payable"))

        # SSC
        ssc_no = str(get_val("ssc_request_number")).strip() if get_val("ssc_request_number") is not None else None
        rec["ssc_request_number"] = ssc_no
        rec["ssc_upload_date"] = format_date_safe(get_val("ssc_upload_date"))
        rec["hard_copy_received_date"] = format_date_safe(get_val("hard_copy_received_date"))
        rec["hard_copy_sent_date"] = format_date_safe(get_val("hard_copy_sent_date"))

        # Payment raw
        due_raw = get_val("payment_due_days")
        rec["payment_due_days_raw"] = str(due_raw).strip() if due_raw is not None else None

        proc_status = str(get_val("invoice_processing_status")).strip() if get_val("invoice_processing_status") is not None else None
        rec["invoice_processing_status"] = proc_status

        rec["processed"] = get_val("processed")
        rec["inprocess"] = get_val("inprocess")

        # Deductions
        rec["remark_closure_date"] = format_date_safe(get_val("remark_closure_date"))
        rec["kz_number"] = str(get_val("kz_number")).strip() if get_val("kz_number") is not None else None
        rec["kz_date"] = format_date_safe(get_val("kz_date"))
        rec["rt"] = clean_amount(get_val("rt"))
        rec["ld"] = clean_amount(get_val("ld"))
        rec["ka"] = clean_amount(get_val("ka"))
        rec["kd"] = clean_amount(get_val("kd"))
        rec["tds"] = clean_amount(get_val("tds"))

        rec["is_duplicate"] = bool(dup_mask.iloc[idx])

        records.append(rec)

    meta = {
        "filename": filename,
        "sheet_name": sheet_name,
        "total_records": len(records),
        "total_gross_value": sum(float(r.get("gross_invoice_total") or 0.0) for r in records),
        "total_amount_payable": sum(float(r.get("total_amount_payable") or 0.0) for r in records),
    }

    return records, meta
